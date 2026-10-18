import io
import re

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT_NAME = "Avenir Next LT Pro"
BODY_PT = 13.5  # Thunderbird "Large" (18 px)
HEADING_PT = 16
LINK_COLOR = "0563C1"
LOOKING_AHEAD = "Looking ahead..."
TABLE_INDENT_TWIPS = 720  # one default tab stop (0.5 in)

# Inline Markdown, in priority order: [text](url), **bold**, *italic* / _italic_, bare URL.
_INLINE = re.compile(
    r"\[(?P<ltext>[^\]]+)\]\((?P<lurl>https?://[^\s)]+)\)"
    r"|\*\*(?P<bold>.+?)\*\*"
    r"|(?<![\w*])\*(?P<ital>[^\s*](?:.*?[^\s*])?)\*(?![\w*])"
    r"|(?<!\w)_(?P<ital2>[^\s_](?:.*?[^\s_])?)_(?!\w)"
    r"|(?P<url>https?://[^\s<>()]*[^\s<>().,;:!?'\"])"
)
_HEADING = re.compile(r"^#{1,6}\s+(.*)$")
_BULLET = re.compile(r"^[-*+]\s+(.*)$")
_NUMBERED = re.compile(r"^\d+[.)]\s+(.*)$")
_RULE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})$")


def build_docx(email_text: str) -> bytes:
    """Convert the Claude-generated email text into a formatted .docx file."""
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = FONT_NAME
    # Also cover East Asian / complex-script text so every run uses the font.
    r_fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), FONT_NAME)
    r_fonts.set(qn("w:cs"), FONT_NAME)
    style.font.size = Pt(BODY_PT)

    date_rows: list[tuple[str, str]] = []
    body_lines: list[str] = []

    for line in email_text.strip().split("\n"):
        if line.startswith("DATE_ROW:"):
            parts = line.replace("DATE_ROW:", "").strip().split("|")
            if len(parts) == 2:
                date_rows.append((parts[0].strip(), parts[1].strip()))
        else:
            body_lines.append(line)

    for line in body_lines:
        stripped = line.strip()
        if stripped == "UPCOMING_TABLE":
            _insert_table(doc, date_rows)
        elif stripped == "" or _RULE.match(stripped) or _is_looking_ahead(stripped):
            # A standalone "Looking ahead" line is re-added directly above the table.
            # Paragraph spacing comes from space_after; no blank paragraphs.
            continue
        elif m := _HEADING.match(stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, m.group(1), bold=True, size=HEADING_PT)
        elif m := _BULLET.match(stripped):
            add_inline(doc.add_paragraph(style="List Bullet"), m.group(1))
        elif m := _NUMBERED.match(stripped):
            add_inline(doc.add_paragraph(style="List Number"), m.group(1))
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            add_inline(p, stripped)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _is_looking_ahead(line: str) -> bool:
    return line.lower().rstrip(":.… ").strip("*_ ") == "looking ahead"


def add_inline(
    paragraph, text: str, bold: bool = False, italic: bool = False, size: float = BODY_PT
):
    """Append *text* to *paragraph*, rendering inline Markdown and URLs as real formatting."""
    pos = 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            _add_run(paragraph, text[pos : m.start()], bold, italic, size)
        if m.group("lurl"):
            _add_hyperlink(paragraph, m.group("ltext"), m.group("lurl"), bold, italic, size)
        elif m.group("url"):
            _add_hyperlink(paragraph, m.group("url"), m.group("url"), bold, italic, size)
        elif m.group("bold") is not None:
            add_inline(paragraph, m.group("bold"), True, italic, size)
        else:
            add_inline(paragraph, m.group("ital") or m.group("ital2"), bold, True, size)
        pos = m.end()
    if pos < len(text):
        _add_run(paragraph, text[pos:], bold, italic, size)


def _add_run(paragraph, text: str, bold: bool, italic: bool, size: float):
    run = paragraph.add_run(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None
    return run


def _add_hyperlink(paragraph, text: str, url: str, bold: bool, italic: bool, size: float) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = _add_run(paragraph, text, bold, italic, size)
    run.font.color.rgb = RGBColor.from_string(LINK_COLOR)
    run.font.underline = True
    # Move the run inside the hyperlink element so Word treats it as clickable.
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _set_cell_bg(cell, hex_color: str) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def _insert_table(doc: Document, date_rows: list[tuple[str, str]]) -> None:
    if not date_rows:
        return

    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)
    heading.paragraph_format.keep_with_next = True
    add_inline(heading, LOOKING_AHEAD, bold=True)

    table = doc.add_table(rows=1 + len(date_rows), cols=2)
    table.style = "Table Grid"
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_TWIPS))
    tbl_ind.set(qn("w:type"), "dxa")
    table._tbl.tblPr.append(tbl_ind)

    hdr = table.rows[0]
    for i, text in enumerate(["Event", "Date & Time"]):
        cell = hdr.cells[i]
        _set_cell_bg(cell, "2E5D2E")
        cell.paragraphs[0].clear()
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.name = FONT_NAME
        run.font.size = Pt(BODY_PT)

    for idx, (event, date_str) in enumerate(date_rows):
        row = table.rows[idx + 1]
        fill = "F2F2F2" if idx % 2 == 0 else "FFFFFF"
        for i, text in enumerate([event, date_str]):
            cell = row.cells[i]
            _set_cell_bg(cell, fill)
            cell.paragraphs[0].clear()
            add_inline(cell.paragraphs[0], text)
