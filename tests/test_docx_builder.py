import io

from docx import Document

from pcc_bulletin.docx_builder import build_docx

FULL_TEXT = """\
Dear Plainfield Christian Church Friends,

This Sunday Steve Higgs brings a powerful message from Romans 8.

Rapha Support Ministry meets Tuesday at 7pm.

See the bulletin for details: https://pccmi.org/bulletin
UPCOMING_TABLE
DATE_ROW: Fall Potluck | Sun, Oct 25 @ 12:00 pm
DATE_ROW: Men's Breakfast | Sat, Oct 31 @ 8:00 am
"""


def _load(docx_bytes: bytes) -> Document:
    return Document(io.BytesIO(docx_bytes))


def test_returns_bytes():
    result = build_docx(FULL_TEXT)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_table_created_with_header():
    doc = _load(build_docx(FULL_TEXT))
    assert len(doc.tables) == 1
    hdr = doc.tables[0].rows[0]
    assert hdr.cells[0].text == "Event"
    assert hdr.cells[1].text == "Date & Time"


def test_table_rows_populated():
    doc = _load(build_docx(FULL_TEXT))
    table = doc.tables[0]
    assert table.rows[1].cells[0].text == "Fall Potluck"
    assert table.rows[1].cells[1].text == "Sun, Oct 25 @ 12:00 pm"
    assert table.rows[2].cells[0].text == "Men's Breakfast"
    assert table.rows[2].cells[1].text == "Sat, Oct 31 @ 8:00 am"


def test_no_table_when_no_rows():
    text = "Dear friends,\n\nSome content.\nUPCOMING_TABLE\n"
    doc = _load(build_docx(text))
    assert len(doc.tables) == 0


def test_malformed_date_row_skipped():
    # DATE_ROW with no pipe delimiter should be silently dropped
    text = "Dear friends,\nUPCOMING_TABLE\nDATE_ROW: No pipe here\nDATE_ROW: Event | Date\n"
    doc = _load(build_docx(text))
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[1].cells[0].text == "Event"


def test_looking_ahead_is_bold():
    text = "Looking ahead:\nSome item.\nUPCOMING_TABLE\nDATE_ROW: E | D\n"
    doc = _load(build_docx(text))
    la_paragraphs = [p for p in doc.paragraphs if p.text == "Looking ahead:"]
    assert len(la_paragraphs) == 1
    assert la_paragraphs[0].runs[0].bold


def test_normal_paragraph_not_bold():
    text = "Dear friends,\nUPCOMING_TABLE\n"
    doc = _load(build_docx(text))
    para = [p for p in doc.paragraphs if p.text == "Dear friends,"]
    assert len(para) == 1
    assert not para[0].runs[0].bold


def test_alternating_row_colors():
    # Build a doc with 3 data rows to hit both colour branches (idx 0, 1, 2)
    text = (
        "UPCOMING_TABLE\n"
        "DATE_ROW: Event A | Mon, Oct 19 @ 9:00 am\n"
        "DATE_ROW: Event B | Tue, Oct 20 @ 10:00 am\n"
        "DATE_ROW: Event C | Wed, Oct 21 @ 11:00 am\n"
    )
    result = build_docx(text)
    # Verify it round-trips without error; colour values are in XML only
    doc = _load(result)
    assert len(doc.tables[0].rows) == 4  # 1 header + 3 data


def test_empty_lines_become_paragraphs():
    text = "Line one.\n\nLine two.\nUPCOMING_TABLE\n"
    doc = _load(build_docx(text))
    texts = [p.text for p in doc.paragraphs]
    assert "Line one." in texts
    assert "" in texts
    assert "Line two." in texts
